import os
import uuid
import json
from typing import Dict, Any, List, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.database import get_db
from app.config import settings, get_wib_now, get_wib_str, format_wib_display

STORAGE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "storage", "exports")
os.makedirs(STORAGE_DIR, exist_ok=True)

class ReportService:
    @staticmethod
    def get_report_summary(user_id: str, start_date: Optional[str] = None, end_date: Optional[str] = None) -> Dict[str, Any]:
        # Validasi Rentang Tanggal
        if start_date and end_date:
            if start_date > end_date:
                raise ValueError("Tanggal mulai laporan tidak boleh lebih besar dari tanggal akhir.")

        conn = get_db()
        cursor = conn.cursor()

        query_income = "SELECT COALESCE(SUM(amount), 0.0) FROM transactions WHERE user_id = ? AND type = 'income'"
        query_expense = "SELECT COALESCE(SUM(amount), 0.0) FROM transactions WHERE user_id = ? AND type = 'expense'"
        query_transfer = "SELECT COALESCE(SUM(amount), 0.0), COUNT(id) FROM transactions WHERE user_id = ? AND type = 'transfer'"

        params_inc = [user_id]
        params_exp = [user_id]
        params_trf = [user_id]

        if start_date:
            query_income += " AND transaction_date >= ?"
            query_expense += " AND transaction_date >= ?"
            query_transfer += " AND transaction_date >= ?"
            params_inc.append(f"{start_date} 00:00:00")
            params_exp.append(f"{start_date} 00:00:00")
            params_trf.append(f"{start_date} 00:00:00")
        if end_date:
            query_income += " AND transaction_date <= ?"
            query_expense += " AND transaction_date <= ?"
            query_transfer += " AND transaction_date <= ?"
            params_inc.append(f"{end_date} 23:59:59")
            params_exp.append(f"{end_date} 23:59:59")
            params_trf.append(f"{end_date} 23:59:59")

        cursor.execute(query_income, params_inc)
        total_income = float(cursor.fetchone()[0])

        cursor.execute(query_expense, params_exp)
        total_expense = float(cursor.fetchone()[0])

        cursor.execute(query_transfer, params_trf)
        row_trf = cursor.fetchone()
        total_transfer_volume = float(row_trf[0])
        total_transfer_count = int(row_trf[1])

        net_cashflow = total_income - total_expense

        # Breakdown Seluruh Kategori Pengeluaran beserta Persentase
        query_cat = """
        SELECT c.id, c.name, c.icon, c.color, COALESCE(SUM(t.amount), 0.0) as total_spent, COUNT(t.id) as transaction_count
        FROM categories c
        JOIN transactions t ON t.category_id = c.id AND t.user_id = c.user_id AND t.type = 'expense'
        WHERE c.user_id = ?
        """
        params_cat = [user_id]
        if start_date:
            query_cat += " AND t.transaction_date >= ?"
            params_cat.append(f"{start_date} 00:00:00")
        if end_date:
            query_cat += " AND t.transaction_date <= ?"
            params_cat.append(f"{end_date} 23:59:59")

        query_cat += " GROUP BY c.id ORDER BY total_spent DESC"
        cursor.execute(query_cat, params_cat)
        cat_rows = cursor.fetchall()

        category_breakdown = []
        for r in cat_rows:
            spent = float(r["total_spent"])
            pct = round((spent / total_expense * 100), 2) if total_expense > 0 else 0.0
            category_breakdown.append({
                "category_id": r["id"],
                "name": r["name"],
                "icon": r["icon"],
                "color": r["color"] or "#FF6B00",
                "total_spent": spent,
                "percentage": pct,
                "transaction_count": int(r["transaction_count"])
            })

        top_categories = category_breakdown[:5]

        conn.close()
        return {
            "total_income": total_income,
            "total_expense": total_expense,
            "net_cashflow": net_cashflow,
            "savings_rate": round((net_cashflow / total_income * 100), 1) if total_income > 0 else 0.0,
            "category_breakdown": category_breakdown,
            "top_categories": top_categories,
            "transfer_summary": {
                "total_volume": total_transfer_volume,
                "count": total_transfer_count,
                "note": "Transfer internal tidak dihitung sebagai pemasukan atau pengeluaran kas."
            }
        }


    @staticmethod
    def get_filtered_transactions(
        user_id: str,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        category_id: Optional[str] = None,
        wallet_id: Optional[str] = None,
        tx_type: Optional[str] = None,
        search_query: Optional[str] = None,
        sort_by: str = "transaction_date",
        sort_order: str = "DESC",
        page: int = 1,
        page_size: int = 15
    ) -> Dict[str, Any]:
        conn = get_db()
        cursor = conn.cursor()

        # Check plan
        cursor.execute("SELECT plan FROM profiles WHERE id = ?", (user_id,))
        p = cursor.fetchone()
        plan = p["plan"] if p else "starter"

        # Starter plan limitation: only monthly summary, no advanced filtering
        if plan == "starter":
            # For Starter, enforce current month only and no category/search filter
            now = get_wib_now()
            start_date = now.strftime("%Y-%m-01")
            category_id = None
            wallet_id = None
            search_query = None

        where_clauses = ["t.user_id = ?"]
        params = [user_id]

        if start_date:
            where_clauses.append("t.transaction_date >= ?")
            params.append(f"{start_date} 00:00:00")
        if end_date:
            where_clauses.append("t.transaction_date <= ?")
            params.append(f"{end_date} 23:59:59")
        if category_id:
            where_clauses.append("t.category_id = ?")
            params.append(category_id)
        if wallet_id:
            where_clauses.append("t.wallet_id = ?")
            params.append(wallet_id)
        if tx_type:
            where_clauses.append("t.type = ?")
            params.append(tx_type)
        if search_query:
            where_clauses.append("t.note LIKE ?")
            params.append(f"%{search_query.strip()}%")

        where_sql = " AND ".join(where_clauses)

        # Count total
        cursor.execute(f"SELECT COUNT(*) FROM transactions t WHERE {where_sql}", params)
        total_records = cursor.fetchone()[0]

        # Valid columns for sorting
        allowed_sorts = {
            "transaction_date": "t.transaction_date",
            "amount": "t.amount",
            "category": "c.name",
            "wallet": "w.name"
        }
        order_col = allowed_sorts.get(sort_by, "t.transaction_date")
        order_dir = "ASC" if sort_order.upper() == "ASC" else "DESC"

        offset = (page - 1) * page_size
        query = f"""
        SELECT t.*, w.name as wallet_name, c.name as category_name, c.icon as category_icon, c.color as category_color
        FROM transactions t
        LEFT JOIN wallets w ON w.id = t.wallet_id
        LEFT JOIN categories c ON c.id = t.category_id
        WHERE {where_sql}
        ORDER BY {order_col} {order_dir}
        LIMIT ? OFFSET ?
        """
        cursor.execute(query, params + [page_size, offset])
        rows = cursor.fetchall()
        conn.close()

        items = []
        for r in rows:
            d = dict(r)
            d["receipt_items"] = json.loads(d.get("receipt_items") or "[]")
            items.append(d)

        return {
            "transactions": items,
            "total_records": total_records,
            "page": page,
            "page_size": page_size,
            "total_pages": (total_records + page_size - 1) // page_size if page_size > 0 else 1,
            "plan": plan
        }

    @staticmethod
    def generate_excel_report(user_id: str, start_date: Optional[str] = None, end_date: Optional[str] = None) -> str:
        # Check plan
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT full_name, plan FROM profiles WHERE id = ?", (user_id,))
        p = cursor.fetchone()
        if not p or p["plan"] != "pro":
            conn.close()
            raise ValueError("Fitur Export Laporan Excel eksklusif untuk pengguna paket Pro.")

        user_name = p["full_name"]

        # Fetch all transactions in range
        where_clauses = ["t.user_id = ?"]
        params = [user_id]
        if start_date:
            where_clauses.append("t.transaction_date >= ?")
            params.append(f"{start_date} 00:00:00")
        if end_date:
            where_clauses.append("t.transaction_date <= ?")
            params.append(f"{end_date} 23:59:59")

        cursor.execute(f"""
        SELECT t.transaction_date, t.type, t.amount, t.note, t.source, t.receipt_items,
               w.name as wallet_name, c.name as category_name
        FROM transactions t
        LEFT JOIN wallets w ON w.id = t.wallet_id
        LEFT JOIN categories c ON c.id = t.category_id
        WHERE {" AND ".join(where_clauses)}
        ORDER BY t.transaction_date ASC
        """, params)
        rows = cursor.fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan TataDana"

        # Theme Colors: Orange (#FF6B00)
        orange_fill = PatternFill(start_color="FF6B00", end_color="FF6B00", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="FF6B00")
        sub_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Title
        ws.cell(row=1, column=1, value="TataDana — Laporan Keuangan Pribadi").font = title_font
        period_text = f"Periode: {start_date or 'Awal'} s/d {end_date or 'Sekarang'} | Pemilik: {user_name} | Generated: {format_wib_display()}"
        ws.cell(row=2, column=1, value=period_text).font = sub_font

        # Headers
        headers = ["No", "Tanggal (WIB)", "Jenis", "Kategori", "Dompet", "Nominal (IDR)", "Catatan", "Sumber", "Rincian Item Struk"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.fill = orange_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 8] else "left", vertical="center")

        total_income = 0.0
        total_expense = 0.0

        for idx, r in enumerate(rows, 1):
            row_idx = 4 + idx
            amount = float(r["amount"])
            tx_type = "Pemasukan" if r["type"] == "income" else ("Pengeluaran" if r["type"] == "expense" else "Transfer")
            if r["type"] == "income":
                total_income += amount
            elif r["type"] == "expense":
                total_expense += amount

            items_str = ""
            try:
                items = json.loads(r["receipt_items"] or "[]")
                if items:
                    items_str = ", ".join([f"{it.get('name')}: Rp{it.get('price', 0):,}" for it in items])
            except Exception:
                pass

            ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=r["transaction_date"][:16])
            ws.cell(row=row_idx, column=3, value=tx_type)
            ws.cell(row=row_idx, column=4, value=r["category_name"] or "-")
            ws.cell(row=row_idx, column=5, value=r["wallet_name"] or "-")
            
            num_cell = ws.cell(row=row_idx, column=6, value=amount)
            num_cell.number_format = '#,##0'
            num_cell.alignment = Alignment(horizontal="right")

            ws.cell(row=row_idx, column=7, value=r["note"] or "-")
            ws.cell(row=row_idx, column=8, value=r["source"] or "web")
            ws.cell(row=row_idx, column=9, value=items_str)

            for c in range(1, 10):
                ws.cell(row=row_idx, column=c).border = thin_border

        # Summary Row
        summary_row = 5 + len(rows)
        ws.cell(row=summary_row, column=5, value="Total Pemasukan").font = bold_font
        c_inc = ws.cell(row=summary_row, column=6, value=total_income)
        c_inc.number_format = '#,##0'
        c_inc.font = bold_font

        ws.cell(row=summary_row+1, column=5, value="Total Pengeluaran").font = bold_font
        c_exp = ws.cell(row=summary_row+1, column=6, value=total_expense)
        c_exp.number_format = '#,##0'
        c_exp.font = bold_font

        ws.cell(row=summary_row+2, column=5, value="Net Cashflow").font = bold_font
        c_net = ws.cell(row=summary_row+2, column=6, value=total_income - total_expense)
        c_net.number_format = '#,##0'
        c_net.font = bold_font

        # Column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        file_id = f"TataDana_Laporan_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}.xlsx"
        file_path = os.path.join(STORAGE_DIR, file_id)
        wb.save(file_path)

        # Save to Export History
        now_wib = get_wib_str()
        cursor.execute("""
        INSERT INTO export_history (id, user_id, file_name, file_type, file_path, storage_url, period_start, period_end, download_count, created_at)
        VALUES (?, ?, ?, 'xlsx', ?, ?, ?, ?, 0, ?)
        """, (str(uuid.uuid4()), user_id, file_id, file_path, f"/api/reports/download/{file_id}", start_date, end_date, now_wib))
        conn.commit()
        conn.close()

        return file_id

    @staticmethod
    def generate_pdf_report(user_id: str, start_date: Optional[str] = None, end_date: Optional[str] = None) -> str:
        # Check plan
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT full_name, plan FROM profiles WHERE id = ?", (user_id,))
        p = cursor.fetchone()
        if not p or p["plan"] != "pro":
            conn.close()
            raise ValueError("Fitur Export Laporan PDF eksklusif untuk pengguna paket Pro.")

        user_name = p["full_name"]

        # Fetch transactions
        where_clauses = ["t.user_id = ?"]
        params = [user_id]
        if start_date:
            where_clauses.append("t.transaction_date >= ?")
            params.append(f"{start_date} 00:00:00")
        if end_date:
            where_clauses.append("t.transaction_date <= ?")
            params.append(f"{end_date} 23:59:59")

        cursor.execute(f"""
        SELECT t.transaction_date, t.type, t.amount, t.note,
               w.name as wallet_name, c.name as category_name
        FROM transactions t
        LEFT JOIN wallets w ON w.id = t.wallet_id
        LEFT JOIN categories c ON c.id = t.category_id
        WHERE {" AND ".join(where_clauses)}
        ORDER BY t.transaction_date DESC
        LIMIT 100
        """, params)
        rows = cursor.fetchall()

        file_id = f"TataDana_Laporan_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}.pdf"
        file_path = os.path.join(STORAGE_DIR, file_id)

        doc = SimpleDocTemplate(file_path, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()
        elements = []

        # Header Title
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=20,
            textColor=colors.HexColor('#FF6B00'),
            spaceAfter=4
        )
        sub_style = ParagraphStyle(
            'SubStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=15
        )

        elements.append(Paragraph("<b>TataDana</b> — Laporan Keuangan", title_style))
        elements.append(Paragraph(f"Nama Pengguna: <b>{user_name}</b> | Periode: {start_date or 'Awal'} s/d {end_date or 'Sekarang'} | Dicetak: {format_wib_display()}", sub_style))

        # Calculate Totals
        total_income = sum(float(r["amount"]) for r in rows if r["type"] == "income")
        total_expense = sum(float(r["amount"]) for r in rows if r["type"] == "expense")
        net = total_income - total_expense

        # Summary Box Table
        summary_data = [
            ["Total Pemasukan", "Total Pengeluaran", "Net Cashflow"],
            [f"Rp {total_income:,.0f}", f"Rp {total_expense:,.0f}", f"Rp {net:,.0f}"]
        ]
        summary_table = Table(summary_data, colWidths=[175, 175, 175])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FFF7ED')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#C2410C')),
            ('TEXTCOLOR', (0, 1), (0, 1), colors.HexColor('#059669')),
            ('TEXTCOLOR', (1, 1), (1, 1), colors.HexColor('#DC2626')),
            ('TEXTCOLOR', (2, 1), (2, 1), colors.HexColor('#2563EB')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 1), (-1, 1), 13),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#FDBA74')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#FED7AA')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 15))

        # Transaction Table
        table_data = [["Tanggal (WIB)", "Jenis", "Kategori", "Dompet", "Nominal (IDR)", "Catatan"]]
        for r in rows:
            t_type = "Pemasukan" if r["type"] == "income" else ("Pengeluaran" if r["type"] == "expense" else "Transfer")
            note = (r["note"] or "-")[:25]
            table_data.append([
                r["transaction_date"][:16],
                t_type,
                r["category_name"] or "-",
                r["wallet_name"] or "-",
                f"Rp {float(r['amount']):,.0f}",
                note
            ])

        tx_table = Table(table_data, colWidths=[85, 75, 105, 95, 85, 80])
        tx_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FF6B00')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (4, 1), (4, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
        ]))
        elements.append(tx_table)

        doc.build(elements)

        # Save to Export History
        now_wib = get_wib_str()
        cursor.execute("""
        INSERT INTO export_history (id, user_id, file_name, file_type, file_path, storage_url, period_start, period_end, download_count, created_at)
        VALUES (?, ?, ?, 'pdf', ?, ?, ?, ?, 0, ?)
        """, (str(uuid.uuid4()), user_id, file_id, file_path, f"/api/reports/download/{file_id}", start_date, end_date, now_wib))
        conn.commit()
        conn.close()

        return file_id

    @staticmethod
    def get_export_history(user_id: str) -> List[Dict[str, Any]]:
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("""
        SELECT id, file_name, file_type, storage_url, period_start, period_end, download_count, created_at
        FROM export_history
        WHERE user_id = ?
        ORDER BY created_at DESC
        """, (user_id,))
        rows = cursor.fetchall()
        conn.close()
        return [dict(r) for r in rows]
